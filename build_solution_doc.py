from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path


OUT = Path("AI_NAS素材理解与文案分镜匹配系统_设计方案_v0.1.docx")
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 100, 112)
LIGHT_FILL = "F2F4F7"
BLUE_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"


def set_font(run, name="Noto Sans CJK SC", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers, rows, widths_dxa, header_fill=LIGHT_FILL, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    set_repeat_table_header(table.rows[0])
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(str(text))
        set_font(r, size=font_size, bold=True, color=DARK_BLUE)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_font(r, size=font_size)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_font(r)
    return p


def new_numbering_id(doc):
    numbering = doc.part.numbering_part.element
    style = doc.styles["List Number"]._element
    style_num_ids = style.xpath("./w:pPr/w:numPr/w:numId")
    if not style_num_ids:
        raise RuntimeError("List Number style has no numbering definition")
    base_num_id = int(style_num_ids[0].get(qn("w:val")))
    base_nums = numbering.xpath(f'./w:num[@w:numId="{base_num_id}"]')
    if not base_nums:
        raise RuntimeError("List Number numbering instance is missing")
    abstract_num_id = base_nums[0].find(qn("w:abstractNumId")).get(qn("w:val"))
    existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    num_id = max(existing_ids, default=0) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_num_id)
    num.append(abstract)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def add_number(doc, text, num_id):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id
    r = p.add_run(text)
    set_font(r)
    return p


def add_callout(doc, label, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run(label + "  ")
    set_font(r1, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_font(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# Preset: standard_business_brief; header pattern: memo_masthead.
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Noto Sans CJK SC"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans CJK SC")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans CJK SC")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
normal.font.size = Pt(11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for style_name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
]:
    st = styles[style_name]
    st.font.name = "Noto Sans CJK SC"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans CJK SC")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans CJK SC")
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = True
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for list_name in ("List Bullet", "List Bullet 2", "List Number"):
    st = styles[list_name]
    st.font.name = "Noto Sans CJK SC"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Noto Sans CJK SC")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Noto Sans CJK SC")
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.167

# Header/footer.
hp = section.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = hp.add_run("AI NAS 素材理解与文案分镜匹配系统  |  技术设计方案")
set_font(hr, size=9, color=MUTED)
fp = section.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = fp.add_run("内部方案 · v0.1 · 2026-07-13")
set_font(fr, size=9, color=MUTED)

# Masthead.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("技术设计方案")
set_font(r, size=11, color=BLUE, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run("AI NAS 素材理解与文案分镜匹配系统")
set_font(r, size=24, color=RGBColor(0, 0, 0), bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
r = p.add_run("基于 SMB 素材库、JoyAI 连续视频理解与文案—分镜—素材匹配")
set_font(r, size=13, color=MUTED)

metadata = [
    ("文档版本", "v0.1"),
    ("日期", "2026-07-13"),
    ("部署基线", "内部 AI 计算节点（硬件型号与实际网络信息已脱敏）"),
    ("产品边界", "不自动剪辑；提供素材理解、搜索、分镜建议、时间码与 NAS 地址"),
]
for label, value in metadata:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(label + "：")
    set_font(r1, bold=True)
    r2 = p.add_run(value)
    set_font(r2)

add_callout(doc, "核心结论", "系统自动发现 SMB 项目素材，但由用户手动确认后启动 AI 分析；JoyAI 负责连续片段理解，自适应视频采样负责控制输入规模并保留时间码。用户粘贴文案后，系统生成分镜表，匹配 NAS 中的具体素材片段，并提供网页预览与 SMB 地址。")

doc.add_heading("1. 项目目标与边界", level=1)
add_body(doc, "本系统面向个人创作者或小型内容团队，将统一存储在 NAS 项目目录中的照片和视频转化为可理解、可搜索、可复用的素材库。系统的主要价值不是替代剪辑师，而是压缩“浏览原片—理解内容—寻找画面—记录文件名和时间码—制作分镜表”的人工链路。")

doc.add_heading("1.1 核心目标", level=2)
for item in [
    "从指定 SMB 项目目录刷新素材清单，并准确识别新增、删除、移动、重命名和内容更新。",
    "对用户确认的素材执行自适应采样和 JoyAI 连续视频理解，生成片段级描述、摘要、关键词及时间码。",
    "支持中文自然语言搜索，并直接返回命中的素材、片段时间码、网页预览和 SMB 地址。",
    "支持粘贴完整文案，自动拆解为画面需求和分镜，匹配现有素材并形成可编辑的分镜表。",
    "保留人工选择权：用户可替换候选素材、锁定素材、标记缺口并导出结果。",
]:
    add_bullet(doc, item)

doc.add_heading("1.2 明确不包含", level=2)
for item in [
    "不自动生成最终剪辑成片，不直接决定剪辑节奏或最终取舍。",
    "第一版不做人脸实名识别、复杂团队权限、外网分享相册或云端协作审片。",
    "第一版不强制生成 Premiere、DaVinci Resolve 或 Final Cut Pro 工程；优先导出 DOCX/Excel/CSV/JSON 分镜数据。",
    "不让浏览器直接访问 SMB 文件；预览统一由网页服务提供代理流或低码率预览。",
]:
    add_bullet(doc, item)

doc.add_heading("2. 用户工作流", level=1)
workflow = [
    "摄影师将现场素材上传至 NAS 项目目录，并完成废片删除、重命名和文件夹整理。",
    "用户在网页中选择项目并点击“刷新素材清单”；系统只读取目录和元数据，不运行 AI。",
    "网页显示视频数量、照片数量、总容量及与上次分析相比的变化。",
    "用户点击“开始 AI 分析”，系统为本次任务生成素材清单快照，并仅处理新增或内容已变化的素材。",
    "分析完成后，用户可通过自然语言、文件名、日期、设备、场景和关键词搜索素材。",
    "用户粘贴文案，系统拆解分镜需求，为每个镜号推荐具体素材及时间码。",
    "用户在网页中预览、替换、锁定或标记素材缺口，最后导出分镜表。",
]
workflow_num_id = new_numbering_id(doc)
for step in workflow:
    add_number(doc, step, workflow_num_id)

doc.add_heading("3. 总体系统架构", level=1)
add_table(
    doc,
    ["层级", "职责", "主要组件"],
    [
        ("存储层", "保存项目原始素材；所有服务节点统一通过 SMB 挂载", "NAS / SMB 共享"),
        ("发现层", "刷新目录、识别变化、生成待分析清单", "目录扫描器、快速指纹、项目快照"),
        ("媒体处理层", "读取元数据、转码预览、自适应采样并保留时间码", "FFmpeg、FFprobe、ExifTool"),
        ("视觉理解层", "理解连续画面变化，输出事件、片段摘要与关键词", "JoyAI-VL-Interaction 8B、可替换 VLM 接口"),
        ("索引层", "保存素材、片段、地址、文本与向量", "PostgreSQL/SQLite、pgvector/Qdrant"),
        ("匹配层", "拆解文案，形成画面需求并匹配素材片段", "语言模型、混合检索、规则打分"),
        ("应用层", "项目管理、搜索、预览、分镜编辑和导出", "局域网 Web 应用"),
    ],
    [1400, 4300, 3660],
)

doc.add_heading("3.1 主处理链路", level=2)
main_flow_num_id = new_numbering_id(doc)
for step in [
    "SMB 项目素材刷新与变化检测",
    "用户手动确认并创建分析快照",
    "视频元数据读取、代理预览和自适应采样",
    "JoyAI 按时间顺序理解连续画面",
    "生成片段描述、整条摘要、关键词和向量",
    "网页搜索或文案分镜匹配",
    "返回命中时间码、网页预览及 SMB 地址",
]:
    add_number(doc, step, main_flow_num_id)

doc.add_heading("4. SMB 与素材地址设计", level=1)
add_body(doc, "所有机器统一通过内部 SMB 挂载访问素材。业务数据库不以某台机器的本地挂载路径作为素材身份，也不向用户显示本地物理路径。")

doc.add_heading("4.1 地址模型", level=2)
add_table(
    doc,
    ["字段", "示例", "用途"],
    [
        ("share_id", "media_nas", "定位配置中的 NAS 共享根"),
        ("relative_path", "川西项目/机位A/DJI_0214.MP4", "与共享根组合为最新地址"),
        ("asset_id", "7842", "系统内部稳定素材编号"),
        ("SMB 地址", "smb://192.0.2.20/素材库/川西项目/机位A/DJI_0214.MP4", "剪辑电脑定位或复制原素材"),
        ("网页预览", "http://ai-media.local/assets/7842?t=41.2", "浏览器从命中时间码开始预览"),
    ],
    [1600, 4700, 3060],
)
add_callout(doc, "设计原则", "分镜只保存 asset_id 与入出点时间码。素材重命名或移动后，系统更新相对路径；分镜引用和网页预览链接保持稳定，导出时自动生成最新 SMB 地址。")

doc.add_heading("4.2 素材身份与变化处理", level=2)
add_table(
    doc,
    ["操作", "判断", "处理"],
    [
        ("删除废片", "原位置消失且 NAS 正常", "标记 missing，经确认后 deleted；取消未开始任务"),
        ("重命名/移动", "快速内容指纹一致", "更新 relative_path，复用分析结果"),
        ("覆盖同名文件", "路径相同但指纹变化", "建立新版本并重新分析"),
        ("复制文件", "内容相同、位置不同", "复用视觉结果并记录新位置"),
        ("SMB 临时断线", "共享整体不可用或大量文件同时消失", "暂停同步，不批量删除数据库记录"),
    ],
    [1700, 3300, 4360],
)

doc.add_heading("5. 分析触发与任务状态", level=1)
add_body(doc, "第一版采用“自动发现、手动分析”。摄影师可以先在 NAS 中自由整理素材，避免对上传中的文件和即将删除的废片消耗推理资源。")
add_table(
    doc,
    ["状态", "含义", "允许的操作"],
    [
        ("整理中", "目录可变化，系统尚未运行 AI", "刷新素材清单"),
        ("待确认", "已识别变化并显示统计", "查看变化、开始分析"),
        ("分析中", "按快照处理素材", "查看进度、暂停、取消未开始任务"),
        ("可检索", "本次索引完成", "搜索、粘贴文案、制作分镜"),
        ("素材有变化", "分析后又出现新增、删除或更新", "仅分析新增和变更"),
    ],
    [1700, 4100, 3560],
)

doc.add_heading("6. 自适应视频采样与 JoyAI 理解", level=1)
add_body(doc, "JoyAI 不能直接逐帧吞入几十或上百 GB 的原视频。自适应采样不是独立的最终功能，而是视频理解前的必要输入控制层：它负责选择有代表性的连续画面并保留准确时间码；JoyAI 负责理解这些画面合起来发生了什么。")

doc.add_heading("6.1 推荐采样策略", level=2)
for item in [
    "按固定间隔进行基础采样，初始建议每 2～5 秒一帧。",
    "检测明显转场，在转场前后补充代表帧。",
    "检测固定机位中的内容变化，例如人物入场、拿起物品或观众动作。",
    "对高度相似的连续帧去重，避免重复占用模型上下文。",
    "按 30～60 秒形成分析窗口；窗口内保持连续会话，窗口间生成阶段摘要。",
    "对疑似重要或变化密集的片段进行二次加密采样。",
]:
    add_bullet(doc, item)

doc.add_heading("6.2 JoyAI 输出契约", level=2)
add_table(
    doc,
    ["输出", "说明"],
    [
        ("segment_description", "片段内发生的事件，强调人物、物体、动作和场景变化"),
        ("start/end", "片段入点和出点，单位为秒，同时可格式化为时间码"),
        ("keywords", "场景、人物、动作、物体、景别、光线和地点类关键词"),
        ("asset_summary", "整条素材的时间顺序摘要"),
        ("quality_flags", "模糊、遮挡、过曝、欠曝、剧烈抖动等辅助标记"),
        ("embedding_text", "用于向量检索的规范化文本"),
    ],
    [2600, 6760],
)
add_callout(doc, "模型策略", "JoyAI 作为首选视频理解模型；视觉分析层必须保留可替换接口，并使用 Qwen3-VL 等模型作为 GPU 运行时兼容性基准或备用后端。")

doc.add_heading("7. 搜索与素材定位", level=1)
add_body(doc, "检索采用精确检索与语义检索的混合方式。文件名、相对路径、日期、设备、时长和人工标签使用结构化查询；“主持人在舞台展示产品”“观众鼓掌”等描述使用文本向量检索。结果按内容、动作、场景、时间和质量综合排序。")
add_table(
    doc,
    ["结果字段", "示例"],
    [
        ("素材", "A0038.MP4"),
        ("命中内容", "工作人员从车辆后备箱取出摄影器材"),
        ("时间码", "02:14–02:27"),
        ("SMB 地址", "smb://nas/项目/机位A/A0038.MP4"),
        ("网页预览", "/assets/842?t=134"),
        ("候选操作", "预览、复制地址、锁定到分镜、查看相似素材"),
    ],
    [2200, 7160],
)

doc.add_heading("8. 文案—分镜—素材匹配", level=1)
add_body(doc, "用户可在网页中粘贴完整口播、旁白或脚本文案。语言模型先将文本拆解为叙事单元和画面需求，匹配引擎再从素材库中寻找满足条件的片段。模型不得凭文件名编造结果；所有素材地址和时间码必须来自索引数据库。")

doc.add_heading("8.1 文案拆解字段", level=2)
for item in [
    "原始文案、分镜意图与预计时长",
    "主体、场景、动作、物体和地点",
    "建议景别、镜头类型、光线与情绪",
    "必须满足条件与可选增强条件",
    "素材是否允许重复、是否要求人物或地点一致",
]:
    add_bullet(doc, item)

doc.add_heading("8.2 匹配排序建议", level=2)
add_table(
    doc,
    ["维度", "建议权重", "说明"],
    [
        ("语义内容", "40%", "文案画面需求与片段描述的整体一致性"),
        ("人物/物体/动作", "20%", "关键主体与行为是否真实出现"),
        ("地点/场景", "15%", "环境、时间和地点是否一致"),
        ("景别/镜头类型", "10%", "大全景、特写、航拍、行车视角等"),
        ("画面质量", "10%", "清晰度、曝光、抖动、遮挡"),
        ("重复惩罚", "5%", "避免连续分镜反复使用同一素材"),
    ],
    [2600, 1800, 4960],
)

doc.add_heading("8.3 分镜表示例", level=2)
add_table(
    doc,
    ["镜号", "文案", "建议画面", "素材/时间码", "交付地址"],
    [
        ("01", "清晨，我们从成都出发", "清晨道路、车辆出发", "A0032.MP4\n00:12–00:19", "SMB 地址\n网页预览"),
        ("02", "穿过连绵的草原", "行车视角或航拍草原", "DJI_0185.MP4\n01:23–01:31", "SMB 地址\n网页预览"),
        ("03", "最终抵达雪山脚下", "雪山大全景、车辆驶入", "DJI_0214.MP4\n00:41–00:50", "SMB 地址\n网页预览"),
        ("04", "同行的人开始整理装备", "人物从后备箱取装备", "C0048.MP4\n02:06–02:15", "SMB 地址\n网页预览"),
    ],
    [700, 2000, 2500, 2200, 1960],
    font_size=8.8,
)
add_callout(doc, "缺口处理", "没有合适素材时必须明确显示“素材缺口”，并说明缺少的主体、动作或场景，不得为了填满分镜而强行匹配低相关素材。")

doc.add_heading("9. 网页功能设计", level=1)
add_table(
    doc,
    ["页面", "主要能力"],
    [
        ("项目页", "选择 SMB 项目、刷新素材清单、查看变化与分析进度"),
        ("素材浏览页", "缩略图网格、故事板、悬停预览、结构化筛选"),
        ("搜索页", "自然语言查询、候选排序、命中时间码、SMB 地址"),
        ("素材详情页", "代理播放、关键帧、AI 摘要、关键词、前后片段和相似素材"),
        ("文案工作台", "粘贴文案、拆分分镜、匹配候选、替换与锁定"),
        ("任务中心", "等待、运行、失败、暂停、重试及资源占用"),
    ],
    [2200, 7160],
)

doc.add_heading("10. 部署基线与硬件判断", level=1)
add_body(doc, "2026-07-13 已对应用计算节点进行只读盘点。归档版只保留设备角色与能力边界，厂商、型号、主机名、精确容量和网络接口均已脱敏。")
add_table(
    doc,
    ["项目", "实测状态", "方案影响"],
    [
        ("节点角色", "应用计算节点", "第一阶段开发和验证节点"),
        ("CPU", "多核通用处理器", "适合扫描、FFmpeg、服务与并发任务"),
        ("GPU", "本地推理加速器", "需验证 JoyAI 与 GPU 运行时兼容性"),
        ("内存", "满足 8B 级模型常驻需求", "同时为缓存和索引预留容量"),
        ("系统", "Linux LTS", "版本按实际部署环境选择"),
        ("GPU 运行时", "厂商与版本已脱敏", "部署时确认驱动、设备代理与 PyTorch"),
        ("存储", "本地高速存储", "用于模型、数据库、关键帧和代理预览"),
        ("网络", "独立局域网接口", "正式接入 NAS 前验证有线 SMB 吞吐"),
    ],
    [1700, 3100, 4560],
)

doc.add_heading("10.1 JoyAI 迁移验证门槛", level=2)
validation_num_id = new_numbering_id(doc)
for step in [
    "确认 GPU 运行时能识别本地推理设备，并记录运行时、驱动和 PyTorch 版本。",
    "以 Transformers 或 vLLM 的兼容后端加载 JoyAI-VL-Interaction-Preview 8B。",
    "使用真实现场图片验证单帧中文描述。",
    "输入连续 5～10 帧验证事件变化理解及会话状态。",
    "用 20 条真实视频完成采样、描述、索引和搜索闭环。",
    "记录单帧耗时、每小时视频处理速度、内存峰值、温度和持续功耗。",
]:
    add_number(doc, step, validation_num_id)

doc.add_heading("11. 数据模型概要", level=1)
add_table(
    doc,
    ["实体", "关键字段"],
    [
        ("Project", "project_id、name、share_id、root_relative_path、analysis_state"),
        ("Asset", "asset_id、project_id、relative_path、content_fingerprint、media_metadata、status"),
        ("Segment", "segment_id、asset_id、start_sec、end_sec、description、keywords、quality_flags"),
        ("Preview", "asset_id、proxy_key、thumbnail_key、storyboard_key"),
        ("AnalysisRun", "run_id、snapshot、model_version、sampling_policy、progress、errors"),
        ("ScriptProject", "script_id、source_text、version、status"),
        ("Shot", "shot_id、script_id、narration、visual_intent、duration、locked_candidate"),
        ("ShotCandidate", "shot_id、segment_id、score、reason、rank、selected"),
    ],
    [2200, 7160],
)

doc.add_heading("12. 非功能要求", level=1)
for item in [
    "安全：SMB 凭据只保存在服务器密钥存储或环境配置中，不进入数据库和网页链接。",
    "稳定性：SMB 断线时暂停同步；批量文件消失不得触发批量删除。",
    "幂等性：重复刷新与重复提交不得造成同一素材被重复分析。",
    "可追溯：每条描述保存模型版本、采样策略和分析批次，便于重新分析或对比。",
    "性能：网页优先读取本地代理文件和索引，不反复从 NAS 解码原视频。",
    "可替换性：视觉模型、文本模型、向量库和数据库通过明确接口解耦。",
]:
    add_bullet(doc, item)

doc.add_heading("13. 实施阶段", level=1)
add_table(
    doc,
    ["阶段", "范围", "完成标志"],
    [
        ("阶段 0：硬件验证", "应用计算节点 GPU 运行时/JoyAI 兼容性和真实素材小样", "连续帧理解可用，形成吞吐与资源报告"),
        ("阶段 1：素材底座", "SMB 项目、刷新、快照、预览、自适应采样、AI 索引", "可搜索并定位具体素材和时间码"),
        ("阶段 2：文案工作台", "文案拆解、分镜需求、素材候选、锁定与导出", "可从文案生成可编辑分镜表"),
        ("阶段 3：质量提升", "景别、人物一致性、重复控制、缺口分析", "匹配准确率达到业务可用标准"),
        ("阶段 4：工程集成", "可选 XML/EDL、多人权限、定时增量", "进入稳定生产工作流"),
    ],
    [1900, 4100, 3360],
)

doc.add_heading("14. MVP 验收标准", level=1)
for item in [
    "一个约 100GB 的已整理项目可以刷新清单，并由用户手动启动分析。",
    "系统能处理文件重命名和移动，保留既有分析结果及稳定网页预览链接。",
    "至少支持 MP4、MOV、JPG、PNG 的元数据、缩略图和网页预览。",
    "用户输入五类中文自然语言描述，可在数秒内返回相关素材及大致正确的片段时间码。",
    "用户粘贴一段真实文案后，系统可生成分镜需求，为多数镜号提供至少三个候选片段。",
    "分镜表可人工修改、锁定和标记缺口，并导出素材文件名、时间码、SMB 地址和网页预览。",
    "NAS 断线、文件删除和分析失败均有明确状态，且不会破坏既有索引。",
]:
    add_bullet(doc, item)

doc.add_heading("15. 风险与待确认事项", level=1)
add_table(
    doc,
    ["事项", "当前判断", "下一步"],
    [
        ("JoyAI 运行时", "计算资源满足预估，软件兼容性需实测", "在应用计算节点运行最小加载与连续帧测试"),
        ("网络吞吐", "无线连接不适合长期读取大量 4K 素材", "启用有线网络并测 SMB 吞吐"),
        ("代理容量", "本地盘可能被代理视频占满", "确定代理码率、保留周期和清理策略"),
        ("匹配质量", "依赖片段切分、描述质量和规则组合", "建立真实文案与素材的评测集"),
        ("项目规范", "不同摄影师目录结构可能差异较大", "定义项目根、机位和已确认目录的最小约定"),
    ],
    [2200, 3900, 3260],
)

add_callout(doc, "建议的立即下一步", "先在应用计算节点上完成只读软件环境盘点和 JoyAI 最小迁移验证；验证通过后，再实现 SMB 项目扫描、代理预览和索引数据库。这样能最早暴露 GPU 运行时兼容性风险，避免先完成网页后才发现核心模型无法稳定运行。")

# Document metadata.
doc.core_properties.title = "AI NAS素材理解与文案分镜匹配系统设计方案"
doc.core_properties.subject = "SMB素材库、JoyAI视频理解、自然语言检索与文案分镜匹配"
doc.core_properties.author = "Codex"
doc.core_properties.keywords = "NAS, SMB, JoyAI, 视频理解, 素材检索, 文案分镜"

doc.save(OUT)
print(OUT.resolve())
